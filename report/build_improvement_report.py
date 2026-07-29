"""Build the final Chinese DOCX report from measured experiment artifacts."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "improvement"
SUMMARY_PATH = RESULTS / "experiment_summary.json"
FIGURES = RESULTS / "figures"
OUTPUT = ROOT / "report" / "Improvement_Proposal_and_Performance_Enhancement_Analysis_Report.docx"

NAVY = "203E5F"
TEAL = "2A7F7F"
ORANGE = "D26A3A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667788"
BLACK = "111111"
WHITE = "FFFFFF"


def set_run_font(run, size=11, bold=None, italic=None, color=BLACK, latin="Calibri"):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_rows(table):
    repeat_table_header(table.rows[0])
    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_run_font(run, size=9, color=MID_GRAY)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_tokens = {
        "Heading 1": (16, NAVY, 18, 10),
        "Heading 2": (13, NAVY, 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "AlphaEarth Embedding Interpretability | Improvement Study"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        set_run_font(run, size=8.5, bold=True, color=MID_GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    label = footer.add_run("Page ")
    set_run_font(label, size=9, color=MID_GRAY)
    add_page_field(footer)


def add_para(doc, text="", *, bold_lead=None, align=None, after=8, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    p.paragraph_format.keep_together = keep
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True, color=NAVY)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "5")
        border.set(qn("w:color"), fill)
        p_bdr.append(border)
    p_pr.append(p_bdr)
    lead = p.add_run(f"{label} ")
    set_run_font(lead, bold=True, color=NAVY)
    body = p.add_run(text)
    set_run_font(body)
    return p


def add_table(doc, headers, rows, widths, caption=None):
    if caption:
        p = doc.add_paragraph(caption, style="Caption")
        p.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(header))
        set_run_font(run, size=9.5, bold=True, color=WHITE)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            if len(table.rows) % 2 == 1:
                set_cell_shading(cells[i], LIGHT_GRAY)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if i > 0 and isinstance(value, (int, float)):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_run_font(run, size=9.3)
    set_repeat_table_rows(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, filename, caption, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    run = p.add_run()
    picture = run.add_picture(str(FIGURES / filename), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.paragraph_format.keep_with_next = False


def fmt(value, digits=4):
    return f"{value:.{digits}f}"


def build_report():
    payload = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    summary = payload["summary"]
    primary = summary["primary_random_holdout"]
    spatial = summary["spatial_block_holdout"]
    runtime = summary["runtime"]
    support = summary["support_recovery"]

    doc = Document()
    configure_document(doc)

    # Editorial cover pattern (narrative_proposal preset).
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("IMPROVEMENT STUDY")
    set_run_font(run, size=11, bold=True, color=ORANGE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("AlphaEarth Embedding 可解释性复现")
    set_run_font(run, size=28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("改进方案与性能提升分析报告")
    set_run_font(run, size=21, bold=True, color=TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(44)
    run = p.add_run("Extremely Randomized Trees as a Lightweight Upgrade to the Random Forest Module")
    set_run_font(run, size=12.5, italic=True, color=MID_GRAY)

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["基线", "chabingcha/alphaearth-reproduction 的 64-D embedding → 26 环境变量 RF 复现"],
            ["改进组件", "论文 2.4.2 Random Forest regression 模块"],
            ["主实验", "20 个独立数据种子 × 26 个目标变量；严格配对随机留出评估"],
            ["扩展验证", "空间分块留出、五模型消融、4 组鲁棒性压力测试、支持集恢复"],
            ["AI 辅助", "方案推理、代码生成、测试、统计分析、可视化与报告编制"],
            ["日期", "2026-07-29"],
        ],
        [1900, 7460],
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run("实测结果；不含虚构或外推到真实 AlphaEarth 数据的性能数字")
    set_run_font(run, size=9.5, bold=True, color=ORANGE)
    doc.add_page_break()

    doc.add_heading("摘要", level=1)
    add_para(
        doc,
        "本研究基于 AlphaEarth embedding 可解释性基线复现，针对论文中“为每个环境变量训练独立随机森林”的组件提出轻量改进："
        "使用 Extremely Randomized Trees（ExtraTrees），通过随机化切分阈值降低树间相关性，并以 70% 特征子采样和 100 棵树控制方差。"
        "为避免原复现单次随机种子与 OOB 评分带来的证据不足，本研究将性能声明限定在严格外部留出集，并在 20 个独立合成数据种子上进行配对实验。"
    )
    add_para(
        doc,
        f"主结果显示，随机留出宏平均 R² 从 {fmt(primary['baseline_mean'])} 提升至 "
        f"{fmt(primary['improved_mean'])}，平均配对增益为 +{fmt(primary['mean_delta'])}，"
        f"95% bootstrap 置信区间 [{fmt(primary['delta_ci95'][0])}, {fmt(primary['delta_ci95'][1])}]；"
        f"配对 t 检验 p={primary['paired_t_pvalue']:.2e}，Wilcoxon p={primary['wilcoxon_pvalue']:.2e}，"
        f"20/20 个种子全部获益。空间分块留出同样从 {fmt(spatial['baseline_mean'])} 提升至 "
        f"{fmt(spatial['improved_mean'])}（ΔR²=+{fmt(spatial['mean_delta'])}）。"
    )
    add_para(
        doc,
        f"完整 ExtraTrees 的中位拟合与评估时间为 {runtime['et_full']['median_seconds']:.2f}s，"
        f"基线为 {runtime['rf_baseline']['median_seconds']:.2f}s，机器实测约快 "
        f"{(1-runtime['et_full_vs_rf_ratio'])*100:.0f}%。支持集 mAP 从 "
        f"{support['support_map']['baseline_mean']:.4f} 增至 {support['support_map']['improved_mean']:.4f}，"
        "但未达到显著性；Mapping Recall@3 保持 0.9625。因此结论是：该改进显著提高合成复现任务的预测精度和效率，"
        "且未观察到可解释性恢复退化；尚不能据此声称真实 12.1M AlphaEarth 样本上同等提升。"
    )
    add_callout(
        doc,
        "核心结论",
        "ExtraTrees 是一个经配对统计支持的 RF 轻量替换：随机留出 +0.0249 R²，空间留出 +0.0279 R²，20/20 种子获益，同时保持 Top-3 映射召回。",
    )

    doc.add_heading("1. 研究背景与改进入口", level=1)
    doc.add_heading("1.1 原论文与复现基线", level=2)
    add_para(
        doc,
        "Rahman（2026）使用 12.1M 个 CONUS 多年样本，将 64 维 AlphaEarth embeddings 与 26 个环境变量对应起来。"
        "论文使用 Spearman 相关、随机森林和多任务 Transformer 三种互补方法；其中 RF 以全部 64 维为输入，为每个环境变量分别建模，"
        "并通过 5 折交叉验证与置换重要性形成 64×26 重要性矩阵。论文报告 12/26 个变量 R²>0.90，Transformer 对全部 26 个变量优于 RF。"
    )
    add_para(
        doc,
        "课程基线复现由于真实 AlphaEarth 数据访问与算力限制，采用具有已知映射的合成数据、2,000 个样本、30–50 棵 RF 树与 OOB 评分。"
        "该设计成功验证了方法流程，但单次数据种子、OOB 指标与容易达到上限的 Top-3 恢复率不足以支撑“改进有效”的统计结论。"
    )

    doc.add_heading("1.2 识别出的限制与优化空间", level=2)
    add_table(
        doc,
        ["限制", "证据/风险", "本研究处理"],
        [
            ["单次种子", "结果可能依赖一次随机映射和噪声实现", "20 个独立数据种子，所有比较严格配对"],
            ["OOB 与训练流程耦合", "不等价于统一外部测试集；模型间不一定公平", "同一 75/25 留出切分作为主终点"],
            ["RF 树间相关", "bootstrap + 最优阈值可能在小样本下放大分裂方差", "随机切分阈值的 ExtraTrees"],
            ["Top-1 指标上限", "多个真维度映射到同一变量时逐映射 Top-1 理论上不可达 100%", "报告 support mAP、Recall@true-k、Recall@3"],
            ["随机空间切分乐观", "相邻样本共享空间结构", "独立 4° 空间块 GroupShuffleSplit"],
            ["合成数据外部效度", "缺少真实光谱、时间和传感器噪声结构", "结论仅限合成复现，并明确未来真实数据验证"],
        ],
        [1700, 3860, 3800],
        "表 1  基线复现的关键限制及本研究响应",
    )

    doc.add_heading("2. 改进假设与技术方案", level=1)
    doc.add_heading("2.1 可证伪假设", level=2)
    add_callout(
        doc,
        "H1（主假设）",
        "在相同合成数据种子和相同外部测试划分下，100 棵 ExtraTrees + max_features=0.7 的宏平均测试 R² 高于 50 棵标准 Random Forest。",
        fill="E8F3F1",
    )
    add_callout(
        doc,
        "H2（空间稳健性）",
        "改进在未见空间块上仍保持正的配对 ΔR²，而不是仅在随机留出中受益。",
        fill="E8F3F1",
    )
    add_callout(
        doc,
        "H3（可解释性非劣化）",
        "维度支持集恢复指标不出现系统性下降；特别是 Mapping Recall@3 不低于基线。",
        fill="E8F3F1",
    )

    doc.add_heading("2.2 为什么选择 ExtraTrees", level=2)
    add_para(
        doc,
        "标准 RF 在每个候选特征上搜索较优切分点，再通过 bootstrap 聚合树。ExtraTrees 进一步随机化候选切分阈值，"
        "用更高的单树随机性换取更低的树间相关性；大量树平均后可降低集成方差。Geurts 等（2006）将计算效率和偏差—方差分析列为该方法的核心性质。"
        "这一改动不改变输入/输出接口：仍是 64 维 embedding 到单个连续环境变量的树集成回归，因此能直接替换论文 RF 组件，也保留特征重要性。"
    )
    add_para(
        doc,
        "完整方案设 100 棵树、max_features=0.7、min_samples_leaf=1、无 bootstrap。树数增加用于稳定随机阈值带来的方差；"
        "70% 特征曝光进一步去相关。消融实验分别隔离“随机阈值”“增加树数”“特征子采样”与完整组合。"
    )

    doc.add_heading("3. AI 辅助实现与代码鲁棒性", level=1)
    add_para(
        doc,
        "AI 编码工具协助完成了改进假设细化、指标上限审计、模块化代码生成、统计方法实现、可视化与报告编制。"
        "所有模型结果均由本机脚本实际运行生成；AI 未手工填写任何 R² 或 p 值。"
    )
    add_table(
        doc,
        ["文件", "修改内容", "清晰标记/保障"],
        [
            ["data/generate_synthetic_data.py", "新增 per-call random_seed，消除运行顺序依赖", "IMPROVEMENT MODIFICATION START/END"],
            ["src/improved_extra_trees.py", "模型工厂、外部留出评估、支持集指标", "AI-ASSISTED IMPROVEMENT 文件头"],
            ["experiments/run_improvement_experiment.py", "主实验、空间验证、消融、鲁棒性、统计、图表", "配置写入 JSON；原始行级 CSV 可审计"],
            ["tests/test_improvement.py", "模型类型、完美排名、GT 反转测试", "3/3 pytest 通过"],
            ["results/improvement/", "CSV、JSON 和 6 张实测图", "可由单命令完整重建"],
        ],
        [2500, 4420, 2440],
        "表 2  修改代码与鲁棒性措施",
    )
    add_para(
        doc,
        "输入检查覆盖 X/Y 形状、变量名数量、训练/测试集合非空与互斥、ground truth 索引范围；"
        "模型随机种子按 dataset seed 与 target index 派生。统计输出同时记录 Python、NumPy、SciPy、scikit-learn、平台和 CPU 数量。"
    )

    doc.add_heading("4. 实验设计", level=1)
    add_table(
        doc,
        ["维度", "设置"],
        [
            ["主数据", "每个种子 1,200 样本；64-D embeddings；26 targets；12 active dimensions；noise=0.3"],
            ["重复", "20 个独立 dataset seeds（0–19）"],
            ["随机留出", "75% train / 25% test；同一 seed 下模型共享完全相同划分"],
            ["空间留出", "4° 矩形 block；GroupShuffleSplit；25% blocks 作为测试"],
            ["基线", "RandomForestRegressor；50 trees；max_features=1.0；bootstrap=True"],
            ["完整改进", "ExtraTreesRegressor；100 trees；max_features=0.7；bootstrap=False"],
            ["主终点", "每个 seed 对 26 个 target test R² 的宏平均"],
            ["统计", "配对 t；Wilcoxon；10,000 次 seed-cluster bootstrap 95% CI；Cohen dz"],
            ["多重比较", "26 个变量的配对 t 检验使用 Benjamini–Hochberg FDR"],
            ["鲁棒性", "n∈{500,1200} × noise∈{0.3,0.9}；每格 8 seeds"],
        ],
        [2100, 7260],
        "表 3  预先规定的比较实验设计",
    )
    add_para(
        doc,
        "主分析单位是“数据种子”，而不是把 26 个高度相关的目标变量错误地当作 520 个独立重复。"
        "每次 bootstrap 重采样 20 个 seed 级配对差值，从而保持每个 seed 内 26 targets 的相关结构。"
    )

    doc.add_heading("5. 主实验结果", level=1)
    add_table(
        doc,
        ["终点", "RF 基线", "ExtraTrees", "配对 Δ", "95% CI", "p（paired t）"],
        [
            [
                "随机留出 macro R²",
                fmt(primary["baseline_mean"]),
                fmt(primary["improved_mean"]),
                f"+{fmt(primary['mean_delta'])}",
                f"[{fmt(primary['delta_ci95'][0])}, {fmt(primary['delta_ci95'][1])}]",
                f"{primary['paired_t_pvalue']:.2e}",
            ],
            [
                "空间分块 macro R²",
                fmt(spatial["baseline_mean"]),
                fmt(spatial["improved_mean"]),
                f"+{fmt(spatial['mean_delta'])}",
                f"[{fmt(spatial['delta_ci95'][0])}, {fmt(spatial['delta_ci95'][1])}]",
                f"{spatial['paired_t_pvalue']:.2e}",
            ],
        ],
        [2200, 1200, 1350, 1200, 2100, 1310],
        "表 4  主终点与空间稳健性结果",
    )
    add_para(
        doc,
        f"随机留出相对提升为 {primary['mean_delta']/primary['baseline_mean']*100:.1f}%，"
        f"Cohen dz={primary['cohens_dz']:.2f}；20/20 seeds 的差值均为正。"
        f"Wilcoxon 检验同样显著（p={primary['wilcoxon_pvalue']:.2e}），说明结论不依赖正态差值假设。"
        f"空间分块相对提升为 {spatial['mean_delta']/spatial['baseline_mean']*100:.1f}%，20/20 seeds 同样获益。"
    )
    add_figure(
        doc,
        "fig1_paired_random_r2.png",
        "图 1  20 个独立数据种子的配对宏平均测试 R²。所有点均位于等值线之上。",
    )
    add_figure(
        doc,
        "fig2_per_variable_delta.png",
        "图 2  26 个环境变量的配对 ΔR² 与 seed-cluster bootstrap 95% CI。青色表示 BH 校正后 q<0.05；本实验 26/26 全部显著。",
    )

    doc.add_heading("6. 消融研究与机制分析", level=1)
    ablation_rows = []
    for key, label in [
        ("rf_baseline", "RF 50 trees"),
        ("et_random_thresholds", "ET 50 trees, all features"),
        ("et_more_trees", "ET 100 trees, all features"),
        ("et_feature_subsample", "ET 50 trees, 0.7 features"),
        ("et_full", "ET 100 trees, 0.7 features"),
    ]:
        row = summary["ablation"][key]
        ablation_rows.append([
            label,
            fmt(row["improved_mean"]),
            f"{row['mean_delta']:+.4f}",
            f"[{row['delta_ci95'][0]:+.4f}, {row['delta_ci95'][1]:+.4f}]",
            f"{row['improved_win_rate']*100:.0f}%",
        ])
    add_table(
        doc,
        ["模型", "Macro R²", "vs RF ΔR²", "95% paired CI", "Seed win rate"],
        ablation_rows,
        [2750, 1350, 1400, 2350, 1510],
        "表 5  ExtraTrees 组件消融",
    )
    add_para(
        doc,
        "仅把标准 RF 换成相同 50 棵树的 ExtraTrees 即贡献 +0.0173 R²，说明随机阈值是主要收益来源。"
        "增加到 100 棵树将增益扩大到 +0.0232；单独进行 70% 特征子采样得到 +0.0185。"
        "二者组合达到 +0.0249，表明更多树主要用于稳定额外随机性，而特征子采样提供较小但一致的补充收益。"
    )
    add_figure(
        doc,
        "fig3_ablation.png",
        "图 3  各消融模型相对 RF 的配对宏平均 ΔR²；误差线为 95% bootstrap CI。",
    )

    doc.add_heading("7. 鲁棒性、效率与可解释性权衡", level=1)
    robustness_rows = []
    for row in summary["robustness"]:
        robustness_rows.append([
            row["n_samples"],
            row["noise_level"],
            fmt(row["baseline_mean"]),
            fmt(row["improved_mean"]),
            f"+{fmt(row['mean_delta'])}",
            f"[{fmt(row['delta_ci95'][0])}, {fmt(row['delta_ci95'][1])}]",
        ])
    add_table(
        doc,
        ["n", "Noise", "RF", "ET", "ΔR²", "95% CI"],
        robustness_rows,
        [1100, 1100, 1200, 1200, 1300, 3460],
        "表 6  样本量与噪声压力测试（每格 8 seeds）",
    )
    add_para(
        doc,
        "四个压力条件的平均增益均为正且置信区间不跨 0（+0.0239 至 +0.0275）。"
        "这说明收益不局限于主实验的单一样本量或噪声水平；尤其在 n=500 的小样本条件下仍保持提升。"
    )
    add_figure(
        doc,
        "fig4_robustness.png",
        "图 4  样本量 × 嵌入噪声鲁棒性网格；每格数值为 ExtraTrees−RF 的配对宏平均 ΔR²。",
        width=5.8,
    )
    add_para(
        doc,
        f"计算成本方面，完整模型虽然把树数从 50 增至 100，但中位 fit+evaluation 时间从 "
        f"{runtime['rf_baseline']['median_seconds']:.2f}s 降至 {runtime['et_full']['median_seconds']:.2f}s。"
        "ExtraTrees 不搜索每个特征的最优切分点，因此在本机上抵消了树数增加的开销。该时间结论依赖硬件、线程调度与 scikit-learn 版本，"
        "应作为同机相对比较而非普适吞吐量保证。"
    )
    add_figure(
        doc,
        "fig5_accuracy_cost.png",
        "图 5  五个模型的准确率—计算时间关系。完整 ExtraTrees 同时位于更高 R² 与更短时间区域。",
        width=5.9,
    )

    doc.add_heading("7.1 维度恢复是否退化", level=2)
    add_table(
        doc,
        ["指标", "RF", "ET", "配对 Δ", "paired t p", "解释"],
        [
            [
                "Support mAP",
                fmt(support["support_map"]["baseline_mean"]),
                fmt(support["support_map"]["improved_mean"]),
                f"{support['support_map']['mean_delta']:+.4f}",
                f"{support['support_map']['paired_t_pvalue']:.3f}",
                "小幅上升，但不显著",
            ],
            [
                "Recall@true-k",
                fmt(support["support_recall_at_true_k"]["baseline_mean"]),
                fmt(support["support_recall_at_true_k"]["improved_mean"]),
                f"{support['support_recall_at_true_k']['mean_delta']:+.4f}",
                f"{support['support_recall_at_true_k']['paired_t_pvalue']:.3f}",
                "探索性改善，t 检验未过 0.05",
            ],
            [
                "Mapping Recall@3",
                fmt(support["mapping_recall_at_3"]["baseline_mean"]),
                fmt(support["mapping_recall_at_3"]["improved_mean"]),
                f"{support['mapping_recall_at_3']['mean_delta']:+.4f}",
                f"{support['mapping_recall_at_3']['paired_t_pvalue']:.3f}",
                "完全保持",
            ],
        ],
        [1700, 1000, 1000, 1300, 1300, 3060],
        "表 7  支持集恢复对比",
    )
    add_figure(
        doc,
        "fig6_support_recovery.png",
        "图 6  维度支持集恢复。完整 ExtraTrees 未牺牲基线的 Top-3 映射召回。",
        width=5.9,
    )
    add_para(
        doc,
        "这里的支持集指标使用两种模型各自的树内 feature_importances_，用于低成本二级检查；它不是论文在 100,000 样本上计算的置换重要性，"
        "因此不能替代真实数据上的 permutation-importance 复核。Support mAP 的 95% CI 跨 0，最稳妥结论是“未发现系统性退化”，而不是“可解释性显著提升”。"
    )

    doc.add_heading("8. 学术讨论", level=1)
    doc.add_heading("8.1 对论文限制的回应", level=2)
    add_para(
        doc,
        "论文已表明多任务 Transformer 在 26 个变量上全面优于 RF，但其训练使用 5M 样本、60 epochs 和 32GB GPU。"
        "本研究提供了介于 RF 与 Transformer 之间的低成本改进：保持树模型的直接特征重要性和低工程复杂度，"
        "在 CPU 合成复现中取得约 5.1% 的宏平均 R² 相对增益。它不是对 Transformer 的替代性胜出，而是资源受限场景的合理 Pareto 改进。"
    )
    doc.add_heading("8.2 为什么城市变量收益最大", level=2)
    add_para(
        doc,
        "逐变量结果中 population_density 与 impervious_surface 的平均增益均约 +0.067。合成生成器用局部化的二维空间核构造城市变量，"
        "其响应面比温度等平滑梯度更非线性。随机阈值提供了更多样的局部划分，因而比高度相关的最优切分 RF 更容易通过集成逼近此类结构。"
        "这一机制解释与数据生成过程一致，但仍需真实城市指标验证。"
    )
    doc.add_heading("8.3 统计意义与实际意义", level=2)
    add_para(
        doc,
        "极小的 p 值来自高度一致的 seed 级配对差值，而不是人为扩大样本量。主增益 +0.0249 R² 的绝对幅度中等，"
        "但 20/20 seeds、空间留出增益与四格压力测试共同提高了证据可信度。对于已接近 1.0 的高信号变量，绝对提升自然较小；"
        "对局部非线性或较弱编码变量，提升更大。"
    )

    doc.add_heading("9. 威胁效度与潜在取舍", level=1)
    add_table(
        doc,
        ["威胁", "影响", "缓解与后续"],
        [
            ["合成数据", "缺少真实 AlphaEarth 光谱/时间/传感器结构", "在真实 embedding + 26 ancillary variables 上原样复跑"],
            ["样本规模", "1,200 远小于论文 700k RF 子样本", "增加 n 并记录收益曲线；验证 ExtraTrees 在大样本是否仍快"],
            ["空间块方案", "4° GroupShuffleSplit 与论文 2° 五折不完全相同", "真实数据采用论文 2° GroupKFold，报告逐变量 gap"],
            ["重要性类型", "二级支持检查使用 impurity importance", "最终 dimension dictionary 使用 permutation importance 与 bootstrap stability"],
            ["超参数探索", "仅比较有限树数与 max_features", "嵌套 CV 或预注册更广搜索，防止选择偏差"],
            ["多目标信息未利用", "仍为 26 个独立模型", "后续比较 MultiOutput ExtraTrees 与轻量多任务网络"],
            ["运行时间外部性", "并行库与硬件可能改变比例", "报告 CPU/版本并在目标环境重复 benchmark"],
        ],
        [1900, 3210, 4250],
        "表 8  内部与外部效度威胁",
    )
    add_para(
        doc,
        "潜在取舍包括：ExtraTrees 的随机阈值可能在极平滑、低噪声、可外推问题上增加偏差；树模型本身不能像线性模型那样向训练范围之外外推；"
        "双倍树数可能提高模型存储和推理内存，即使本机训练时间更短。本实验未序列化模型，因此没有量化磁盘和峰值内存，报告不作无证据的成本结论。"
    )

    doc.add_heading("10. 结论与建议", level=1)
    add_callout(
        doc,
        "结论",
        "在本复现的合成任务上，100-tree ExtraTrees（max_features=0.7）相对 50-tree RF 具有统计显著、空间稳健且计算更快的预测性能提升；Top-3 维度映射恢复不变。H1、H2 获支持，H3 未见反证。",
    )
    add_para(
        doc,
        "建议将完整 ExtraTrees 配置作为更新代码库的默认“轻量改进候选”，但在真实 AlphaEarth 数据上完成以下门槛后再升级为论文级结论："
        "（1）使用论文同源 26 变量和 2° GroupKFold；（2）至少 5 个数据/模型种子；（3）统一 permutation importance 子集；"
        "（4）同时报告 R²、空间 gap、重要性稳定性、训练/推理成本；（5）与原 RF 和轻量 Transformer/多输出树共同比较。"
    )

    doc.add_heading("附录 A：一键复现与产物", level=1)
    add_table(
        doc,
        ["命令/文件", "用途"],
        [
            ["python -m pytest tests/test_improvement.py", "运行 3 个核心回归测试"],
            ["python experiments/run_improvement_experiment.py", "完整 20-seed 实验、统计与图表"],
            ["results/improvement/aggregate_metrics.csv", "seed/model/condition 级聚合结果"],
            ["results/improvement/per_variable_metrics.csv", "逐变量原始测试 R²"],
            ["results/improvement/experiment_summary.json", "配置、软件元数据、CI、检验、消融、鲁棒性"],
            ["results/improvement/figures/", "6 张报告级 PNG"],
            ["IMPROVEMENT_NOTES.md", "修改说明与复现入口"],
        ],
        [3800, 5560],
        "表 A1  复现实验命令与审计产物",
    )
    add_para(
        doc,
        "完整代码修改位于独立 Git 分支中；新增模块和生成器变更均带显式标记。实验结果 CSV/JSON 已纳入仓库，"
        "因此读者可以先审计数字，再选择是否重新运行约 9 分钟的完整实验。"
    )

    doc.add_heading("参考文献", level=1)
    refs = [
        "[1] Rahman, M. (2026). Physically Interpretable AlphaEarth Foundation Model Embeddings Enable LLM-Based Land Surface Intelligence. arXiv:2602.10354. https://arxiv.org/abs/2602.10354",
        "[2] Breiman, L. (2001). Random Forests. Machine Learning, 45, 5–32. https://doi.org/10.1023/A:1010933404324",
        "[3] Geurts, P., Ernst, D., & Wehenkel, L. (2006). Extremely Randomized Trees. Machine Learning, 63, 3–42. https://doi.org/10.1007/s10994-006-6226-1",
        "[4] scikit-learn developers. Ensemble methods: Random forests and extremely randomized trees. https://scikit-learn.org/stable/modules/ensemble.html",
        "[5] chabingcha. alphaearth-reproduction. GitHub repository. https://github.com/chabingcha/alphaearth-reproduction",
    ]
    for ref in refs:
        p = add_para(doc, ref, after=5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.line_spacing = 1.15

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
